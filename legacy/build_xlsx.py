"""Формує Excel-таблицю з flagman_raw.json.

Аркуші:
  * «Параметри»     — поля для задання діапазону ваги та довжини;
  * «Усі позиції»   — усі товари (з колонкою «Фільтр»);
  * «Тільки акції»  — лише знижки (теж з колонкою «Фільтр»).

Колонка «Фільтр» повертає 1, якщо товар підпадає під межі з «Параметрів».

Запуск:  uv run python build_xlsx.py
"""

import json
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAW = r"C:\Users\Gleam\OneDrive\Desktop\lures\flagman_raw.json"
OUT = r"C:\Users\Gleam\OneDrive\Desktop\lures\flagman_voblery.xlsx"

HEADERS = [
    "№",
    "Фільтр",
    "Назва",
    "Артикул",
    "Виробник",
    "Тип воблера",
    "Плавучість",
    "Вага, г",
    "Довжина, мм",
    "Заглиблення від, м",
    "Заглиблення до, м",
    "Ціна, грн",
    "Стара ціна, грн",
    "% акції",
    "Рейтинг",
    "Відгуки",
    "Наявність",
    "Посилання",
    "Фото",
]
WIDTHS = {
    1: 6,
    2: 8,
    3: 55,
    4: 16,
    5: 18,
    6: 13,
    7: 12,
    8: 10,
    9: 12,
    10: 17,
    11: 16,
    12: 12,
    13: 15,
    14: 10,
    15: 9,
    16: 9,
    17: 20,
    18: 42,
    19: 42,
}
NCOL = len(HEADERS)

# індекси ключових колонок (1-based)
COL_WEIGHT = HEADERS.index("Вага, г") + 1
COL_LENGTH = HEADERS.index("Довжина, мм") + 1
LETTER_WEIGHT = get_column_letter(COL_WEIGHT)
LETTER_LENGTH = get_column_letter(COL_LENGTH)

# ключові слова в назві/slug -> тип воблера
TYPE_WORDS = [
    ("vib", "Vib/Rattlin"),
    ("rattlin", "Vib/Rattlin"),
    ("crank", "Crank"),
    ("minnow", "Minnow"),
    ("popper", "Popper"),
    ("walker", "Walker"),
    ("swimbait", "Swimbait"),
    ("swim bait", "Swimbait"),
    ("jerk", "Jerkbait"),
    ("cicada", "Cicada"),
    ("fat", "Fat"),
]

# ключові слова в описі -> тип воблера (для моделей без характеристики)
DESC_WORDS = [
    ("міноу", "Minnow"),
    ("мінноу", "Minnow"),
    ("минноу", "Minnow"),
    ("minnow", "Minnow"),
    ("кренк", "Crank"),
    ("крэн", "Crank"),
    ("crank", "Crank"),
    ("шайнер", "Shad"),
    ("шед", "Shad"),
    ("shad", "Shad"),
    ("поппер", "Popper"),
    ("popper", "Popper"),
    ("уокер", "Walker"),
    ("волкер", "Walker"),
    ("walker", "Walker"),
    ("джерк", "Jerkbait"),
    ("jerk", "Jerkbait"),
    ("раттлін", "Vib/Rattlin"),
    ("rattlin", "Vib/Rattlin"),
    ("lipless", "Vib/Rattlin"),
    ("свімбейт", "Swimbait"),
    ("swimbait", "Swimbait"),
    ("цикада", "Cicada"),
    ("cicada", "Cicada"),
    ("фет-кренк", "Fat"),
    ("fat crank", "Fat"),
]

# слово на сайті -> код (поверхневі визначаються за типом: Popper/Walker)
BUOY_MAP = {
    "Плаваючий": "F",
    "Поверхневий": "F",
    "Суспендр": "SP",
    "Тонучий": "S",
    "Швидко тоне": "FS",
    "Повільно тоне": "SS",
    "Повільно cпливає": "SR",
    "Повільно спливає": "SR",
    "Глибоко пірнаючий": "S",
}

d = json.load(open(RAW, encoding="utf-8"))


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_num(s):
    if s is None:
        return None
    m = re.search(r"-?\d+(?:[.,]\d+)?", str(s))
    if not m:
        return None
    return float(m.group().replace(",", "."))


def parse_depth(s):
    """Повертає (від, до) у метрах. 'до X' -> (None, X); одиночне X -> (X, X)."""
    if not s:
        return None, None
    v = s.replace("\u2013", "-").replace("\u2014", "-").replace(",", ".")
    v = re.sub(r"\s+", " ", v).strip().lower()
    nums = [float(x) for x in re.findall(r"\d*\.?\d+", v)]
    if not nums:
        return None, None
    if v.startswith("до "):
        return None, nums[0]
    return min(nums), max(nums)


def char(r, *keywords, exclude=()):
    """Повертає value_ua/value характеристики, у назві якої є будь-яке з keywords."""
    for v in (r.get("characteristics") or {}).values():
        title = (v.get("title_ua") or "") + " " + (v.get("title") or "")
        if any(k in title for k in keywords) and not any(x in title for x in exclude):
            return v.get("value_ua") or v.get("value")
    return None


def buoyancy(r):
    raw = char(r, "лав")
    code = BUOY_MAP.get((raw or "").strip())
    if code:
        return code
    # дозаповнення з назви/моделі
    low = ((r.get("name") or "") + " " + (r.get("title") or "")).lower()
    if "vib" in low:
        return "S"
    if re.search(r"\d\s*sp\b|\bsp[-_ ]", low) or "sp-" in low or "sp " in low:
        return "SP"
    if re.search(r"\d\s*sr\b|\bsr[-_ ]", low) or "sr-" in low:
        return "SR"
    if "bomber" in low and "model a" in low:
        return "F"
    if re.search(r"\d\s*f\b|\d+f-", low):
        return "F"
    return None


def desc_type(r):
    """Тип воблера з опису (перша згадка у тексті)."""
    text = re.sub(r"<[^>]+>", " ", (r.get("description_ua") or r.get("description") or ""))
    text = text.lower()
    best_i, best = None, None
    for kw, val in DESC_WORDS:
        i = text.find(kw)
        if i != -1 and (best_i is None or i < best_i):
            best_i, best = i, val
    return best


def wobbler_type(r):
    """Тип воблера: спершу з характеристики, потім з назви/slug, потім з опису."""
    t = char(r, "Тип воблера")
    if t:
        return t
    low = ((r.get("title") or "") + " " + (r.get("name") or "")).lower()
    for kw, val in TYPE_WORDS:
        if kw in low:
            return val
    return desc_type(r) or ""


def build(r):
    price = num(r.get("price"))
    old = num(r.get("old_price"))
    pct = None
    if price is not None and old and old > price:
        pct = int(round((old - price) / old * 100))
    images = r.get("images") or []
    t = wobbler_type(r) or ""
    # поверхневі (SURF) — це поппери і волкери, а не значення характеристики
    buoy = "SURF" if t in ("Popper", "Walker") else (buoyancy(r) or "")
    depth_from, depth_to = parse_depth(char(r, "аглиб"))
    return {
        "title": r.get("title") or "",
        "sku": r.get("ext_id") or "",
        "producer": (r.get("producer") or {}).get("title") or "",
        "type": t,
        "buoy": buoy,
        "weight": parse_num(char(r, "Вага", exclude=("діапазон",))),
        "length": parse_num(char(r, "Довжина")),
        "depth_from": depth_from,
        "depth_to": depth_to,
        "price": price,
        "old": old,
        "pct": pct,
        "rating": r.get("rating") if r.get("rating") is not None else "",
        "comments": r.get("count_comments") if r.get("count_comments") is not None else "",
        "avail": "В наявності" if str(r.get("sell_status")) == "4" else "Немає в наявності",
        "url": f"https://flagman.ua/{r.get('name')}/p{r.get('id')}",
        "img": images[0] if images else "",
    }


rows = [build(r) for r in d]  # порядок як на сайті (сторінки 1..41)

wb = Workbook()
head_fill = PatternFill("solid", fgColor="0B8BFF")
head_font = Font(bold=True, color="FFFFFF", size=11)
sale_fill = PatternFill("solid", fgColor="FFF3F3")
surf_fill = PatternFill("solid", fgColor="E8F4FF")
match_fill = PatternFill("solid", fgColor="C6EFCE")
input_fill = PatternFill("solid", fgColor="FFF9C4")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")

# ── аркуш «Параметри» ────────────────────────────────────────────────────────
ws_p = wb.active
ws_p.title = "Параметри"
ws_p["A1"] = "Фільтр ваги та довжини"
ws_p["A1"].font = Font(bold=True, size=14, color="0B8BFF")
ws_p["A3"] = "Вага від, г"
ws_p["A4"] = "Вага до, г"
ws_p["A5"] = "Довжина від, мм"
ws_p["A6"] = "Довжина до, мм"
for r in range(3, 7):
    ws_p.cell(row=r, column=1).font = Font(bold=True)
    c = ws_p.cell(row=r, column=2)
    c.fill = input_fill
    c.border = box
    c.number_format = "#,##0.###"
    c.alignment = center
ws_p["A8"] = (
    "Впишіть межі (порожнє поле = без обмеження). Далі на аркуші даних: автофільтр → «Фільтр» = 1."
)
ws_p["A8"].font = Font(italic=True, color="808080")
ws_p["A9"] = "Значення перераховуються автоматично після зміни меж."
ws_p["A9"].font = Font(italic=True, color="808080")
ws_p.column_dimensions["A"].width = 22
ws_p.column_dimensions["B"].width = 14
ws_p.sheet_properties.tabColor = "0B8BFF"


def filter_formula(er):
    """Формула колонки «Фільтр» для рядка er (межі — на аркуші «Параметри»)."""
    return (
        "=IF(AND("
        f'IF(Параметри!$B$3="",TRUE,{LETTER_WEIGHT}{er}>=Параметри!$B$3),'
        f'IF(Параметри!$B$4="",TRUE,{LETTER_WEIGHT}{er}<=Параметри!$B$4),'
        f'IF(Параметри!$B$5="",TRUE,{LETTER_LENGTH}{er}>=Параметри!$B$5),'
        f'IF(Параметри!$B$6="",TRUE,{LETTER_LENGTH}{er}<=Параметри!$B$6)'
        "),1,0)"
    )


def fill_sheet(ws, data):
    ws.append(HEADERS)
    for c in range(1, NCOL + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
    ws.row_dimensions[1].height = 26

    for i, r in enumerate(data, start=1):
        er = i + 1
        ws.cell(row=er, column=1, value=i)
        ws.cell(row=er, column=2, value=filter_formula(er))
        ws.cell(row=er, column=3, value=r["title"])
        ws.cell(row=er, column=4, value=r["sku"])
        ws.cell(row=er, column=5, value=r["producer"])
        ws.cell(row=er, column=6, value=r["type"])
        buoy_cell = ws.cell(row=er, column=7, value=r["buoy"])
        ws.cell(row=er, column=8, value=r["weight"])
        ws.cell(row=er, column=9, value=r["length"])
        ws.cell(row=er, column=10, value=r["depth_from"])
        ws.cell(row=er, column=11, value=r["depth_to"])
        ws.cell(row=er, column=12, value=r["price"])
        ws.cell(row=er, column=13, value=r["old"])
        pct_cell = ws.cell(row=er, column=14, value=r["pct"])
        ws.cell(row=er, column=15, value=r["rating"])
        ws.cell(row=er, column=16, value=r["comments"])
        ws.cell(row=er, column=17, value=r["avail"])
        link = ws.cell(row=er, column=18, value=r["url"])
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single")
        photo = ws.cell(row=er, column=19, value=r["img"])
        if r["img"]:
            photo.hyperlink = r["img"]
            photo.font = Font(color="0563C1", underline="single")

        ws.cell(row=er, column=1).alignment = center
        for col in (COL_WEIGHT, COL_LENGTH, 10, 11):
            ws.cell(row=er, column=col).number_format = "#,##0.###"
        ws.cell(row=er, column=12).number_format = "#,##0.00"
        ws.cell(row=er, column=13).number_format = "#,##0.00"
        for col in (2, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17):
            ws.cell(row=er, column=col).alignment = center

        is_surf = r["buoy"] == "SURF"
        if is_surf:
            buoy_cell.font = Font(bold=True, color="0B5AA6")
            for col in range(1, NCOL + 1):
                ws.cell(row=er, column=col).fill = surf_fill
        if r["pct"]:
            pct_cell.font = Font(bold=True, color="ED1C24")
            if not is_surf:
                for col in range(1, NCOL + 1):
                    ws.cell(row=er, column=col).fill = sale_fill

    for col, w in WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(NCOL)}{len(data) + 1}"
    # підсвітка зеленим, коли «Фільтр» = 1
    ws.conditional_formatting.add(
        f"B2:B{len(data) + 1}",
        CellIsRule(operator="equal", formula=["1"], fill=match_fill),
    )
    return ws


ws1 = wb.create_sheet("Усі позиції")
fill_sheet(ws1, rows)

sales = sorted([r for r in rows if r["pct"]], key=lambda r: -r["pct"])
ws2 = wb.create_sheet("Тільки акції")
fill_sheet(ws2, sales)

wb.active = 0  # відкривати на аркуші «Параметри»
wb.save(OUT)
print(f"Збережено: {OUT}")
print(f"Усі позиції: {len(rows)} | зі знижкою: {len(sales)} | без знижки: {len(rows) - len(sales)}")
print(
    f"Поверхневих (SURF): {sum(1 for r in rows if r['buoy'] == 'SURF')} | без плавучості: {sum(1 for r in rows if not r['buoy'])}"
)
print(
    f"Без типу воблера: {sum(1 for r in rows if not r['type'])} | без заглиблення: {sum(1 for r in rows if r['depth_from'] is None and r['depth_to'] is None)}"
)
print(f"Колонок: {NCOL} (вага={LETTER_WEIGHT}, довжина={LETTER_LENGTH}) | аркушів: {wb.sheetnames}")
