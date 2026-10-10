"""Excel-експорт (openpyxl), узагальнена версія оригінального ``build_xlsx.py``.

Аркуші:
  * «Параметри»   — поля для задання діапазону ваги та довжини;
  * «Усі позиції» — усі товари (з колонкою «Фільтр»);
  * «Тільки акції»— лише знижки.

Колонка «Фільтр» повертає 1, якщо товар підпадає під межі з «Параметрів».
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .canonical import CANONICAL_FIELDS

__all__ = ["build_workbook_bytes"]

PREFIX_HEADERS = ["№", "Фільтр", "Магазин", "Категорія"]
SHEET_PARAMS = "Параметри"
SHEET_ALL = "Усі позиції"
SHEET_SALE = "Тільки акції"

HEAD_FILL = PatternFill("solid", fgColor="0B8BFF")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
SALE_FILL = PatternFill("solid", fgColor="FFF3F3")
SURF_FILL = PatternFill("solid", fgColor="E8F4FF")
MATCH_FILL = PatternFill("solid", fgColor="C6EFCE")
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

NUMBER_FORMAT = "#,##0.###"
MONEY_FORMAT = "#,##0.00"


def _headers() -> list[str]:
    return [*PREFIX_HEADERS, *(field.label for field in CANONICAL_FIELDS)]


def _column_index(field_key: str) -> int:
    index = len(PREFIX_HEADERS)
    for field in CANONICAL_FIELDS:
        index += 1
        if field.key == field_key:
            return index
    raise KeyError(field_key)


def _filter_formula(row: int, weight_letter: str, length_letter: str) -> str:
    return (
        "=IF(AND("
        f'IF({SHEET_PARAMS}!$B$3="",TRUE,{weight_letter}{row}>={SHEET_PARAMS}!$B$3),'
        f'IF({SHEET_PARAMS}!$B$4="",TRUE,{weight_letter}{row}<={SHEET_PARAMS}!$B$4),'
        f'IF({SHEET_PARAMS}!$B$5="",TRUE,{length_letter}{row}>={SHEET_PARAMS}!$B$5),'
        f'IF({SHEET_PARAMS}!$B$6="",TRUE,{length_letter}{row}<={SHEET_PARAMS}!$B$6)'
        "),1,0)"
    )


def _widths() -> dict[int, int]:
    widths: dict[int, int] = {1: 6, 2: 8, 3: 14, 4: 16}
    for field in CANONICAL_FIELDS:
        if field.key == "name":
            widths[_column_index(field.key)] = 55
        elif field.key in ("url", "image_url"):
            widths[_column_index(field.key)] = 42
        else:
            widths[_column_index(field.key)] = 12
    return widths


def _params_sheet(workbook: Workbook) -> None:
    ws = workbook.active
    ws.title = SHEET_PARAMS
    ws["A1"] = "Фільтр ваги та довжини"
    ws["A1"].font = Font(bold=True, size=14, color="0B8BFF")
    ws["A3"] = "Вага від, г"
    ws["A4"] = "Вага до, г"
    ws["A5"] = "Довжина від, мм"
    ws["A6"] = "Довжина до, мм"
    for row in range(3, 7):
        ws.cell(row=row, column=1).font = Font(bold=True)
        cell = ws.cell(row=row, column=2)
        cell.fill = INPUT_FILL
        cell.border = BOX
        cell.number_format = NUMBER_FORMAT
        cell.alignment = CENTER
    ws["A8"] = (
        "Впишіть межі (порожнє поле = без обмеження). "
        "Далі на аркуші даних: автофільтр → «Фільтр» = 1."
    )
    ws["A9"] = "Значення перераховуються автоматично після зміни меж."
    for address in ("A8", "A9"):
        ws[address].font = Font(italic=True, color="808080")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.sheet_properties.tabColor = "0B8BFF"


def _fill_sheet(
    ws: Any, rows: list[dict[str, Any]], weight_letter: str, length_letter: str
) -> None:
    ncols = len(_headers())
    ws.append(_headers())
    for column in range(1, ncols + 1):
        cell = ws.cell(row=1, column=column)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
    ws.row_dimensions[1].height = 26

    weight_col = _column_index("weight_g")
    length_col = _column_index("length_mm")
    depth_from_col = _column_index("depth_from_m")
    depth_to_col = _column_index("depth_to_m")
    buoyancy_col = _column_index("buoyancy")
    price_col = _column_index("price")
    old_price_col = _column_index("old_price")
    discount_col = _column_index("discount_pct")
    rating_col = _column_index("rating")
    comments_col = _column_index("comments")
    availability_col = _column_index("availability")

    for index, row in enumerate(rows, start=1):
        excel_row = index + 1
        ws.cell(row=excel_row, column=1, value=index)
        ws.cell(
            row=excel_row,
            column=2,
            value=_filter_formula(excel_row, weight_letter, length_letter),
        )
        ws.cell(
            row=excel_row,
            column=3,
            value=row.get("store_name") or row.get("store_code") or "",
        )
        ws.cell(row=excel_row, column=4, value=row.get("category_name") or "")

        for column, field in enumerate(CANONICAL_FIELDS, start=len(PREFIX_HEADERS) + 1):
            value = row.get(field.key)
            if field.key in ("url", "image_url") and value:
                cell = ws.cell(row=excel_row, column=column, value=value)
                cell.hyperlink = str(value)
                cell.font = Font(color="0563C1", underline="single")
            else:
                ws.cell(row=excel_row, column=column, value=value)

        ws.cell(row=excel_row, column=1).alignment = CENTER
        for column in (
            weight_col,
            length_col,
            depth_from_col,
            depth_to_col,
            buoyancy_col,
            price_col,
            old_price_col,
            discount_col,
            rating_col,
            comments_col,
            availability_col,
        ):
            ws.cell(row=excel_row, column=column).alignment = CENTER
        for column in (weight_col, length_col, depth_from_col, depth_to_col):
            ws.cell(row=excel_row, column=column).number_format = NUMBER_FORMAT
        ws.cell(row=excel_row, column=price_col).number_format = MONEY_FORMAT
        ws.cell(row=excel_row, column=old_price_col).number_format = MONEY_FORMAT

        is_surf = row.get("buoyancy") == "SURF"
        if is_surf:
            for column in range(1, ncols + 1):
                ws.cell(row=excel_row, column=column).fill = SURF_FILL
        if row.get("discount_pct"):
            ws.cell(row=excel_row, column=discount_col).font = Font(bold=True, color="ED1C24")
            if not is_surf:
                for column in range(1, ncols + 1):
                    ws.cell(row=excel_row, column=column).fill = SALE_FILL

    for column, width in _widths().items():
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "C2"
    last = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last}"
    ws.conditional_formatting.add(
        f"B2:B{last}", CellIsRule(operator="equal", formula=["1"], fill=MATCH_FILL)
    )


def build_workbook_bytes(rows: list[dict[str, Any]]) -> bytes:
    """Побудувати Excel-книгу з канонічних рядків (з DB) і повернути байти."""
    workbook = Workbook()
    _params_sheet(workbook)
    weight_letter = get_column_letter(_column_index("weight_g"))
    length_letter = get_column_letter(_column_index("length_mm"))

    ws_all = workbook.create_sheet(SHEET_ALL)
    _fill_sheet(ws_all, rows, weight_letter, length_letter)

    sales = sorted(
        (row for row in rows if row.get("discount_pct")),
        key=lambda row: int(row.get("discount_pct") or 0),
        reverse=True,
    )
    ws_sale = workbook.create_sheet(SHEET_SALE)
    _fill_sheet(ws_sale, sales, weight_letter, length_letter)

    workbook.active = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
